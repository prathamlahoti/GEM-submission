"""Create six sample-specific Word studies from completed Objectives 1-3 outputs.

Requires pandas, numpy, matplotlib, Pillow, and python-docx. Does not fit models.
Run from the GEM workspace containing analysis_output and GEM_Submission.
"""

from __future__ import annotations

from pathlib import Path
from statistics import median
import csv

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from generate_objectives23_manuscripts import Manuscript, OUTCOMES, pretty

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
MODELS = ("Logistic Regression", "Lasso", "Ridge", "Elastic Net", "Naive Bayes",
          "Decision Tree", "XGBoost", "CatBoost", "Random Forest", "H2O GBM")
SLUG = {m: m.lower().replace(" ", "_") for m in MODELS}
THEMES = {1: "Socio-cognitive traits", 2: "Socio-cognitive traits and the Legatum Prosperity Index",
          3: "Socio-cognitive traits and the Fraser Economic Freedom of the World Index"}
INST = {2: ("prosperity_index_score", "safety__security", "personal_freedom", "governance",
            "social_capital", "investment_environment", "enterprise_conditions",
            "market_access__infrastructure", "economic_quality", "living_conditions", "health",
            "education", "natural_environment"),
        3: ("EFW", "EFW_quartile", "EFW_government", "EFW_legal", "EFW_money",
            "EFW_trade", "EFW_regulation")}


def load(path):
    if not path.exists(): raise FileNotFoundError(path)
    return pd.read_csv(path)


class Study(Manuscript):
    def __init__(self, objective: int, sample: str):
        self.objective, self.sample = objective, sample
        date = "20260927" if objective == 1 else "20260928"
        self.batch = ROOT / "analysis_output" / f"objective{objective}_batch_{date}"
        if not (self.batch / "pipeline_complete.txt").exists():
            raise RuntimeError(f"Objective {objective} batch is incomplete")
        status = load(self.batch / "model_status.csv")
        if len(status) != 120 or not status.status.eq("success").all():
            raise RuntimeError(f"Objective {objective}: model statuses are not 120/120 successful")
        self.casebase = self.batch / f"objective_{objective}" / sample.lower()
        self.dir = HERE / "split_studies" / f"Objective_{objective}" / sample
        self.dir.mkdir(parents=True, exist_ok=True)
        self.figdir = self.dir / "figures"
        self.figdir.mkdir(exist_ok=True)
        self.samples = load(self.batch / "all_sample_counts.csv")
        self.perf = load(self.batch / "all_performance.csv")
        self.comp = load(self.batch / "feature_set_comparison.csv")
        self.doc = Document()
        self._setup()
        for target in OUTCOMES:
            a=(self.casebase/"iv"/target/"heldout_rowids.csv").read_bytes()
            b=(self.casebase/"iv_cv"/target/"heldout_rowids.csv").read_bytes()
            if a!=b: raise RuntimeError(f"Nonmatching held-out records: {objective} {sample} {target}")

    def case(self,target,file,feature="iv_cv"):
        return self.casebase/feature/target/file

    def sample_row(self,target):
        x=self.samples[(self.samples.scope==self.sample.lower())&
                       (self.samples.feature_set=="iv_cv")&
                       (self.samples.outcome==target)]
        return x.iloc[0]

    def perf_rows(self,target,feature="iv_cv"):
        x=self.perf[(self.perf.scope==self.sample.lower())&
                    (self.perf.feature_set==feature)&(self.perf.outcome==target)]
        return x.set_index("model").loc[list(MODELS)].reset_index()

    def best(self,target):
        return self.perf_rows(target).sort_values("auc_pr",ascending=False).iloc[0]

    def imp(self,target):
        x=load(self.case(target,"random_forest_importance.csv"))
        x["share"]=x.importance/x.importance.sum()
        return x.sort_values("share",ascending=False)

    def target_summary(self, fn):
        return "; ".join(f"{t}: {fn(t)}" for t in OUTCOMES)

    def two_paras(self, first, second):
        self.doc.add_heading("Interpretation",level=2)
        self.para(first)
        self.para(second)

    def section(self,title):
        self.doc.add_heading(title,level=1)

    def intro(self):
        scope="India" if self.sample=="India" else "Global"
        self.doc.add_paragraph(f"Objective {self.objective} {scope} results and interpretation",style="Title")
        p=self.doc.add_paragraph()
        p.paragraph_format.space_after=Pt(11)
        r=p.add_run(THEMES[self.objective]);r.italic=True;r.font.size=Pt(11)
        self.section("Abstract")
        lead=("This study evaluates prediction within India" if scope=="India" else
              "This study evaluates prediction in the global sample, including India")
        self.para(f"{lead} for total early-stage entrepreneurial activity (TEA), opportunity-motivated "
                  "entrepreneurship (OME), and necessity-motivated entrepreneurship (NME). "
                  f"Objective {self.objective} combines {THEMES[self.objective].lower()} with ten machine-learning "
                  "classifiers. Independent-variable-only and independent-variable-plus-control specifications "
                  "are compared on identical held-out respondents. Tables and figures use completed saved outputs; "
                  "no classifier was refitted for this manuscript. The evidence supports predictive interpretation, "
                  "not causal attribution.")
        self.section("Study design")
        r=self.sample_row("TEA")
        self.para(f"The analytical sample comprises {int(r.n):,} respondents surveyed across {int(r.years)} years"+
                  (f" and {int(r.countries)} countries." if scope=="Global" else " in India.")+
                  " Three binary outcomes were modeled separately. The ten candidate classifiers were logistic "
                  "regression, lasso, ridge, elastic net, naïve Bayes, decision tree, XGBoost, CatBoost, random "
                  "forest, and H2O gradient boosting. Predictive measures were calculated from the saved held-out "
                  "test predictions, using thresholds selected on validation data.")
        self.para("The principal specification includes the independent variables and the pre-specified controls. "
                  "Its comparison with the independent-variable-only specification uses exactly the same test "
                  "records. Precision-recall area is interpreted against the outcome prevalence, particularly for "
                  "less common positive cases. Receiver operating characteristic area, precision, recall, and the "
                  "F1 score provide complementary views of discrimination and classification.")

    def samples_table(self):
        self.section("Sample composition")
        self.caption("Table 1. Sample size, positive outcome prevalence, and data partition")
        rr=[self.sample_row(t) for t in OUTCOMES]
        self.table(["Outcome","Respondents","Positive cases","Prevalence","Training","Validation","Test"],
                   [[t,f"{int(r.n):,}",f"{int(r.positive):,}",f"{100*r.prevalence:.1f}%",
                     f"{int(r.n_train):,}",f"{int(r.n_validation):,}",f"{int(r.n_test):,}"]
                    for t,r in zip(OUTCOMES,rr)], [.70,1.08,1.10,.85,1.08,1.08,1.07])
        self.note("The two predictor specifications use identical respondent partitions within each target.")
        prev=self.target_summary(lambda t:f"{100*self.sample_row(t).prevalence:.1f}%")
        self.two_paras(
            f"The outcome prevalence is {prev}. Although all three targets use the same respondent pool, their positive classes differ substantially. This matters because a model can achieve high overall accuracy by identifying the more common negative class.",
            f"The held-out sample contains {int(rr[0].n_test):,} observations per outcome. Model comparisons within this document are therefore made on a consistent test partition; results should not be compared with the other geographic manuscript as if the samples were independent, because the global sample includes India.")

    def descriptive_table(self):
        self.doc.add_page_break()
        self.section("Predictor distributions")
        self.caption("Table 2. Descriptive statistics and missingness of model inputs")
        features=load(self.case("TEA","features.csv"))
        stats=load(self.case("TEA","descriptive_statistics.csv")).set_index("variable")
        lines=[]
        for _,f in features.iterrows():
            x=stats.loc[f.variable]
            if x.type=="categorical":
                val=f"{len(str(x.categories).split(';'))} categories"
            elif x.minimum==0 and x.maximum==1:
                val=f"{100*x['mean']:.1f}% coded one"
            else: val=f"{x['mean']:.2f} ({x.sd:.2f})"
            lines.append([pretty(f.variable),"Independent" if f.role=="independent" else "Control",
                          val,f"{100*x.n_missing/x.n:.1f}%"])
        self.table(["Variable","Role","Observed distribution","Missing"],lines,[2.50,1.10,2.21,1.15])
        self.note("Binary variables show the observed percentage coded one. Other numeric variables show mean and standard deviation. Categorical fields show the number of observed categories; missingness is calculated over all eligible respondents.")
        high=sorted(((float(stats.loc[f.variable].n_missing/stats.loc[f.variable].n),pretty(f.variable))
                     for _,f in features.iterrows()),reverse=True)[:3]
        self.two_paras(
            "The table records the distribution of all independent variables and controls used in the principal "
            "specification. The largest missingness percentages occur for "+", ".join(f"{n} ({100*v:.1f}%)" for v,n in high)+". These fields can influence prediction through both their observed values and their recorded missing states.",
            ("Institutional values are shared by respondents within each country and year; respondent totals must not be interpreted as independent institutional observations. " if self.sample=="Global" and self.objective>1 else
             "The national institutional measures are constant within each India survey year, so their variation is primarily temporal. " if self.objective>1 else
             "The socio-cognitive fields capture respondents' reported skills, perceptions, networks, and attitudes. ")+
            "Descriptive differences do not establish adjusted or causal relationships with entrepreneurial outcomes.")

    def missingness_table(self):
        self.section("Entrepreneurial intention response pattern")
        self.caption("Table 3. Missing entrepreneurial intention by observed outcome")
        lines=[]
        for t in OUTCOMES:
            d=load(self.case(t,"intention_missingness_by_outcome.csv"))
            def counts(y,missing):
                z=d[(d.outcome==y)&(d.intention_missing.astype(str).str.upper()==str(missing).upper())]
                return int(z.Freq.iloc[0]) if len(z) else 0
            pos=counts(1,True); ptotal=pos+counts(1,False)
            neg=counts(0,True); ntotal=neg+counts(0,False)
            lines.append([t,f"{ptotal:,}",f"{100*pos/ptotal:.1f}%",f"{ntotal:,}",f"{100*neg/ntotal:.1f}%"])
        self.table(["Outcome","Positive cases","Missing among positives","Negative cases","Missing among negatives"],
                   lines,[.72,1.30,1.65,1.30,1.99])
        self.note("The entrepreneurial-intention field remains a predictor, with its missing state retained under the agreed model specification.")
        self.two_paras(
            "Entrepreneurial intention is missing for every positive case across the three outcomes. The negative-case missingness rates differ by outcome, as shown in the final column. Missingness is therefore strongly associated with outcome status in the observed survey data.",
            "The model can use this response pattern to distinguish positive and negative cases. Consequently, high predictive scores should be interpreted as the performance of the full observed-data pipeline, rather than as evidence that reported intention or institutional conditions alone accurately identify entrepreneurs.")

    def performance_table(self):
        self.section("Held-out model performance")
        self.caption("Table 4. Leading classifier by precision-recall area for each outcome")
        lines=[]
        for t in OUTCOMES:
            x=self.best(t)
            lines.append([t,x.model,f"{x.auc_roc:.3f}",f"{x.auc_pr:.3f}",f"{x.precision:.3f}",
                          f"{x.recall:.3f}",f"{x.f1:.3f}"])
        self.table(["Outcome","Classifier","ROC area","PR area","Precision","Recall","F1"],
                   lines,[.72,1.65,.93,.88,.97,.89,.92])
        self.note("The displayed classifier has the largest saved held-out precision-recall area within each target. All ten classifiers are reported in Supplementary Table S1. ROC denotes receiver operating characteristic; PR denotes precision-recall.")
        bests=self.target_summary(lambda t:f"{self.best(t).model}, precision-recall area {self.best(t).auc_pr:.3f}")
        self.two_paras(
            f"The leading observed models are {bests}. Precision-recall area is the principal ranking measure because the opportunity- and necessity-motivated outcomes have relatively few positive cases. The receiver operating characteristic values provide a complementary assessment of ranking across both classes.",
            "These results were selected descriptively from the held-out metrics and are not evidence that the top classifier is statistically superior to a close competitor. The intention-missingness pattern in Table 3 is present in every fitted model and must be considered when interpreting apparently high discrimination.")

    def confusion_table(self):
        self.section("Classification at selected thresholds")
        self.caption("Table 5. Confusion counts for the leading classifiers")
        lines=[]; desc=[]
        for t in OUTCOMES:
            m=self.best(t).model
            x=load(self.case(t,"confusion_matrices.csv"))
            x=x[x.model==m].iloc[0]
            if int(x.TP+x.FP+x.TN+x.FN)!=int(self.sample_row(t).n_test):
                raise RuntimeError("Confusion counts do not sum to test sample")
            lines.append([t,m,f"{int(x.TP):,}",f"{int(x.FP):,}",f"{int(x.TN):,}",f"{int(x.FN):,}"])
            desc.append(f"{t}: {int(x.FP):,} false positives and {int(x.FN):,} false negatives")
        self.table(["Outcome","Classifier","True positive","False positive","True negative","False negative"],
                   lines,[.75,1.58,1.14,1.17,1.15,1.17])
        self.note("Counts use thresholds selected from validation observations and applied to the held-out test sample. All ten-model confusion results appear in Supplementary Table S2.")
        self.two_paras(
            "The selected classifiers produce "+"; ".join(desc)+". Unlike curve areas, these counts show the concrete false-alarm and missed-case trade-off at the chosen operating points.",
            "The preferred balance of false positives and false negatives depends on the intended use of a classifier. Because no such decision cost was specified, these validation-selected thresholds are reproducible benchmarks rather than policy-optimal decision rules.")

    def controls_table(self):
        self.section("Additional information from controls")
        self.caption("Table 6. Change in precision-recall area after adding controls")
        lines=[]; summaries=[]
        for t in OUTCOMES:
            m=self.best(t).model
            a=self.perf_rows(t,"iv").set_index("model");b=self.perf_rows(t,"iv_cv").set_index("model")
            delta=b.auc_pr-a.auc_pr
            lines.append([t,m,f"{a.loc[m,'auc_pr']:.3f}",f"{b.loc[m,'auc_pr']:.3f}",
                          f"{delta.loc[m]:+.3f}",f"{delta.median():+.3f}"])
            summaries.append(f"{t}: selected-model change {delta.loc[m]:+.3f}, ten-model median {delta.median():+.3f}")
        self.table(["Outcome","Classifier","Predictors only","With controls","Change","Median change"],
                   lines,[.72,1.46,1.34,1.24,1.04,1.16])
        self.note("Change equals the score with controls minus the score with independent variables only, on identical held-out respondents. Full matched comparisons appear in Supplementary Table S3.")
        self.two_paras(
            "The changes are "+"; ".join(summaries)+". The model-level median limits reliance on a single selected classifier and shows whether the control set tends to add predictive information across algorithms.",
            ("Within India, survey year and the national institutional indices share the same temporal variation. " if self.sample=="India" and self.objective>1 else "")+
            "The comparison measures the predictive contribution of the control set as a whole on this test split. It does not identify the separate contribution or causal effect of any individual control.")

    def institutional_table(self):
        if self.objective==1:return
        self.doc.add_page_break()
        self.section("Institutional predictor importance")
        self.caption("Table 7. Leading institutional variables in random-forest importance")
        lines=[]; descriptions=[]
        for t in OUTCOMES:
            x=self.imp(t); z=x[x.variable.isin(INST[self.objective])]
            top=z.iloc[0]
            lines.append([t,pretty(top.variable),f"{100*top.share:.1f}%",f"{100*z.share.sum():.1f}%"])
            descriptions.append(f"{t}: {pretty(top.variable)} ({100*top.share:.1f}% of all importance)")
        self.table(["Outcome","Leading institutional variable","Share of all importance","All institutional fields"],
                   lines,[.75,2.70,1.75,1.76])
        self.note("Shares use all fitted random-forest features as the denominator. The ranking includes raw institutional fields and excludes their missingness indicators, although those indicators remain in the denominator.")
        self.two_paras(
            "The leading institutional fields are "+"; ".join(descriptions)+". The fourth column shows the combined share of all raw institutional variables, allowing the leading pillar to be read in the context of the wider institutional block.",
            ("In India, each score has at most eight annual profiles and is partly confounded with survey year. " if self.sample=="India" else
             "Across countries, multiple institutional pillars are correlated and may substitute for one another in forest splits. ")+
            "Impurity importance describes model usage, not the direction or causal magnitude of an institutional effect.")

    def figure(self,number,title,path,note,p1,p2):
        self.doc.add_page_break()
        self.caption(f"Figure {number}. {title}")
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(2)
        p.add_run().add_picture(str(path),width=Inches(5.7))
        self.note(note)
        self.two_paras(p1,p2)

    def charts(self):
        plt.rcParams.update({"font.family":"DejaVu Serif","font.size":8.5,
                             "axes.spines.top":False,"axes.spines.right":False,
                             "savefig.facecolor":"white"})
        for kind,n,title in (("roc",1,"Receiver operating characteristic curves"),
                              ("pr",2,"Precision-recall curves")):
            fig,axes=plt.subplots(3,1,figsize=(6.4,7.3),sharex=True,sharey=True)
            for ax,t in zip(axes,OUTCOMES):
                metric="auc_roc" if kind=="roc" else "auc_pr"
                best=self.perf_rows(t).nlargest(5,metric)
                for k,(_,r) in enumerate(best.iterrows()):
                    fn=f"models/{SLUG[r.model]}/{kind}_coordinates.csv.gz"
                    d=load(self.case(t,fn))
                    x=d.fpr if kind=="roc" else d.recall
                    y=d.tpr if kind=="roc" else d.precision
                    ix=np.unique(np.linspace(0,len(d)-1,min(1500,len(d)),dtype=int))
                    ax.plot(x.iloc[ix],y.iloc[ix],color=str(.10+.17*k),
                            linestyle=("-","--","-.",":",(0,(5,2,1,2)))[k],linewidth=1.4,
                            label=f"{r.model} ({r[metric]:.3f})")
                if kind=="roc":ax.plot([0,1],[0,1],color=".7",linestyle=":",linewidth=.8)
                else:ax.axhline(self.sample_row(t).prevalence,color=".7",linestyle=":",linewidth=.8)
                ax.set_title(OUTCOMES[t].capitalize(),loc="left",fontweight="bold")
                ax.set_xlim(0,1);ax.set_ylim(0,1);ax.grid(alpha=.12)
                ax.set_ylabel("True-positive rate" if kind=="roc" else "Precision")
                ax.legend(loc="lower right" if kind=="roc" else "upper right",frameon=False,fontsize=6.1)
            axes[-1].set_xlabel("False-positive rate" if kind=="roc" else "Recall")
            fig.subplots_adjust(left=.13,right=.98,top=.98,bottom=.07,hspace=.22)
            path=self.figdir/f"figure_{n}.png";fig.savefig(path,dpi=220,bbox_inches="tight");plt.close(fig)
            vals=self.target_summary(lambda t:
                f"{self.perf_rows(t).nlargest(1,metric).iloc[0].model} "
                f"{self.perf_rows(t).nlargest(1,metric).iloc[0][metric]:.3f}")
            self.figure(n,title,path,
                        "Five strongest models in each panel by the area represented. Dotted line indicates chance discrimination or positive-case prevalence.",
                        f"The largest displayed area values are {vals}. The curves display ranking across all possible thresholds, complementing the single-threshold performance values in Table 4.",
                        "Several curves lie close together, so small visual differences should not be interpreted as established superiority. The outcome-linked intention missingness in Table 3 remains part of these observed predictive results.")
        fig,axes=plt.subplots(3,1,figsize=(6.4,7.3))
        for ax,t in zip(axes,OUTCOMES):
            d=self.imp(t).head(7).iloc[::-1]
            ax.barh([pretty(v) for v in d.variable],100*d.share,color=".35",height=.7)
            ax.set_title(OUTCOMES[t].capitalize(),loc="left",fontweight="bold")
            ax.set_xlabel("Share of total importance (%)");ax.grid(axis="x",alpha=.12)
            ax.tick_params(axis="y",labelsize=7,length=0)
        fig.subplots_adjust(left=.29,right=.98,top=.98,bottom=.07,hspace=.43)
        path=self.figdir/"figure_3.png";fig.savefig(path,dpi=220,bbox_inches="tight");plt.close(fig)
        tops=self.target_summary(lambda t:f"{pretty(self.imp(t).iloc[0].variable)} ({100*self.imp(t).iloc[0].share:.1f}%)")
        self.figure(3,"Random-forest feature importance",path,
                    "The seven largest features in each full-specification random forest are shown as shares of total impurity importance.",
                    f"The leading fitted features are {tops}. Their order differs across outcome definitions, indicating that the forest uses different observed inputs when identifying the three entrepreneurial statuses.",
                    "Importance is a descriptive model diagnostic. It does not reveal the sign of an association or an intervention effect, and the missingness of entrepreneurial intention can contribute to its rank.")
        fig,axes=plt.subplots(3,1,figsize=(6.4,7.3),sharex=True)
        for ax,t in zip(axes,OUTCOMES):
            a=self.perf_rows(t,"iv").set_index("model");b=self.perf_rows(t).set_index("model")
            delta=(b.auc_pr-a.auc_pr).loc[list(MODELS)];y=np.arange(len(MODELS))
            ax.axvline(0,color="black",linewidth=.8)
            ax.hlines(y,0,delta,color=".5",linewidth=1.7)
            ax.plot(delta,y,"o",color=".2",markersize=2.8)
            ax.set_yticks(y,MODELS,fontsize=6.8);ax.invert_yaxis()
            ax.set_title(OUTCOMES[t].capitalize(),loc="left",fontweight="bold")
            ax.grid(axis="x",alpha=.12)
        axes[-1].set_xlabel("Change in precision-recall area after controls")
        fig.subplots_adjust(left=.27,right=.98,top=.98,bottom=.07,hspace=.3)
        path=self.figdir/"figure_4.png";fig.savefig(path,dpi=220,bbox_inches="tight");plt.close(fig)
        changes=self.target_summary(lambda t:f"median {(self.perf_rows(t).set_index('model').auc_pr-self.perf_rows(t,'iv').set_index('model').auc_pr).median():+.3f}")
        self.figure(4,"Change in precision-recall area after adding controls",path,
                    "Each dot is a ten-model paired comparison on exactly the same held-out respondents. Positive values favor the specification with controls.",
                    f"Across model families the changes are {changes}. The chart displays heterogeneity around these medians rather than reducing the comparison to the leading model in Table 6.",
                    "The plotted increment belongs to the full control set, including survey year in the controlled specification. It should not be attributed to any single variable or interpreted causally.")
        if self.objective>1:
            fig,axes=plt.subplots(3,1,figsize=(6.4,7.3))
            for ax,t in zip(axes,OUTCOMES):
                d=self.imp(t);d=d[d.variable.isin(INST[self.objective])].head(6).iloc[::-1]
                ax.barh([pretty(v) for v in d.variable],100*d.share,color=".35",height=.7)
                ax.set_title(OUTCOMES[t].capitalize(),loc="left",fontweight="bold")
                ax.set_xlabel("Share of total importance (%)");ax.grid(axis="x",alpha=.12)
                ax.tick_params(axis="y",labelsize=7,length=0)
            fig.subplots_adjust(left=.3,right=.98,top=.98,bottom=.07,hspace=.43)
            path=self.figdir/"figure_5.png";fig.savefig(path,dpi=220,bbox_inches="tight");plt.close(fig)
            self.figure(5,"Institutional feature importance",path,
                        "The six leading raw institutional fields in each random forest are shown as shares of all fitted-feature importance.",
                        "The three outcome panels make clear whether the institutional ranking is stable across total, opportunity-motivated, and necessity-motivated entrepreneurship. The bar scale uses the whole model as its denominator, not the institutional block alone.",
                        ("India has only eight national annual profiles, and survey year enters the controlled specification. " if self.sample=="India" else
                         "Across countries, index components can be strongly correlated. ")+
                        "The bars describe model use of these fields rather than independent institutional effects.")

    def supplementary(self):
        self.doc.add_page_break();self.section("Supplementary model results")
        self.para("The following tables show all ten candidate classifiers for every outcome. They retain the complete model-level evidence underlying the concise main tables.")
        self.caption("Supplementary Table S1. Full held-out classifier performance")
        lines=[]
        for t in OUTCOMES:
            for _,r in self.perf_rows(t).iterrows():
                lines.append([t,r.model,f"{r.auc_roc:.3f}",f"{r.auc_pr:.3f}",
                              f"{r.precision:.3f}",f"{r.recall:.3f}",f"{r.f1:.3f}"])
        self.table(["Outcome","Classifier","ROC area","PR area","Precision","Recall","F1"],lines,
                   [.7,1.62,.94,.91,.95,.93,.91])
        self.note("All values are from held-out test data for the full independent-variable-plus-control specification.")
        self.caption("Supplementary Table S2. Full held-out confusion matrices")
        lines=[]
        for t in OUTCOMES:
            d=load(self.case(t,"confusion_matrices.csv")).set_index("model")
            for m in MODELS:
                r=d.loc[m]
                lines.append([t,m,*[f"{int(r[k]):,}" for k in ("TP","FP","TN","FN")]])
        self.table(["Outcome","Classifier","True positive","False positive","True negative","False negative"],
                   lines,[.73,1.54,1.17,1.17,1.17,1.18])
        self.note("Every row sums to the test sample size for its outcome; thresholds were selected on validation data.")
        self.caption("Supplementary Table S3. Full matched comparison of predictor specifications")
        lines=[]
        for t in OUTCOMES:
            a=self.perf_rows(t,"iv").set_index("model");b=self.perf_rows(t).set_index("model")
            for m in MODELS:
                lines.append([t,m,f"{a.loc[m,'auc_pr']:.3f}",f"{b.loc[m,'auc_pr']:.3f}",
                              f"{b.loc[m,'auc_pr']-a.loc[m,'auc_pr']:+.3f}"])
        self.table(["Outcome","Classifier","Predictors only","With controls","Change"],lines,
                   [.76,1.70,1.55,1.55,1.40])
        self.note("Each pair was evaluated on identical held-out row identifiers. Positive change favors the specification with controls.")

    def discussion(self):
        self.doc.add_page_break();self.section("Discussion")
        scores=self.target_summary(lambda t:f"{self.best(t).auc_pr:.3f} for {self.best(t).model}")
        self.para(f"Across the three outcomes, the leading precision-recall areas are {scores}. The pattern indicates that "
                  "the saved observed-data specification separates entrepreneurial status from the negative class "
                  "well on the held-out partition. Differences among model families are generally smaller than "
                  "the separation from a prevalence-only benchmark, but no formal uncertainty interval was calculated.")
        self.para("The dominant qualification is the survey structure of entrepreneurial intention: the field is missing "
                  "for every positive case. Retaining this variable followed the agreed specification, but the missing "
                  "state itself offers target-linked information. These discrimination measures should not be interpreted "
                  "as the performance of a prospective survey in which the intention question is asked of every respondent.")
        if self.objective>1:
            self.para(("For India, national institutional scores vary only across the eight survey years, and year is "
                       "included among the controls. Institutional importance therefore reflects predictive use of "
                       "annual variation, without isolating an independent institutional effect." if self.sample=="India" else
                       "In the global sample, institutional measures vary by country and year, but related pillars are "
                       "correlated and the model was fitted for prediction. Importance rankings do not identify a "
                       "policy effect or support causal comparison between institutions."))
        self.section("Conclusion")
        self.para("The study documents outcome-specific predictive performance, the incremental value of the complete "
                  "control set, and the variables used by random forests in this sample. The complete model results "
                  "are retained in the supplementary tables to make the principal comparisons auditable. Substantive "
                  "claims about socio-cognitive or institutional mechanisms would require analyses designed for "
                  "identification rather than this predictive design.")

    def build(self):
        self.intro();self.samples_table();self.descriptive_table();self.missingness_table()
        self.performance_table();self.confusion_table();self.controls_table()
        self.institutional_table();self.charts();self.discussion();self.supplementary()
        path=self.dir/f"Objective_{self.objective}_{self.sample}_Study.docx"
        self.doc.save(path)
        print(path)


if __name__=="__main__":
    for n in (1,2,3):
        for sample in ("India","Global"):
            Study(n,sample).build()
