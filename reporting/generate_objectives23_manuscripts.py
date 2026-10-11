"""Write journal-style Word reports from saved Objective 2 and 3 CSV tables."""

from __future__ import annotations

import csv
from pathlib import Path
from statistics import median

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
OUTCOMES = {"TEA": "total early-stage entrepreneurial activity",
            "OME": "opportunity-motivated entrepreneurship",
            "NME": "necessity-motivated entrepreneurship"}
THEMES = {2: "Legatum Prosperity Index", 3: "Fraser Economic Freedom of the World Index"}


def rows(path):
    with open(path, newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def pick(data, sample, target, model=None):
    result = [r for r in data if r.get("Sample") == sample and r.get("Target") == target]
    if model is not None:
        result = [r for r in result if r.get("Model") == model]
    return result[0] if model is not None else result


def pretty(variable):
    names = {
        "suskill_2015_2022": "Entrepreneurial self-efficacy",
        "fearfail_2015_2022": "Fear of failure",
        "knowing_entrepreneur": "Knowledge of an entrepreneur",
        "entrepreneurial_intention": "Entrepreneurial intention",
        "prosperity_index_score": "Overall prosperity score",
        "EFW": "Overall economic freedom score",
        "EFW_quartile": "Economic freedom quartile",
        "EFW_government": "Government size",
        "EFW_legal": "Legal system and property rights",
        "EFW_money": "Sound money",
        "EFW_trade": "Freedom to trade",
        "EFW_regulation": "Regulation",
    }
    return names.get(variable, variable.replace("__missing", " missing").replace("__", " and ").replace("_", " ").capitalize())


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_border(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def set_cell_margin(cell, top=70, start=65, bottom=70, end=65):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, width in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{side}")
        node = mar.find(tag)
        if node is None:
            node = OxmlElement(f"w:{side}")
            mar.append(node)
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


class Manuscript:
    def __init__(self, objective):
        self.objective = objective
        self.dir = HERE / f"objective{objective}_complete"
        self.data = {n: rows(self.dir / "tables" / f"table_{n}_{suffix}.csv") for n, suffix in (
            (1, "samples"), (2, "descriptives"), (3, "intention_missingness"),
            (4, "performance"), (5, "confusion_matrices"),
            (6, "controls_comparison"), (7, "institutional_importance"))}
        self.doc = Document()
        self._setup()

    def _setup(self):
        sec = self.doc.sections[0]
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.top_margin = Inches(.76)
        sec.bottom_margin = Inches(.72)
        sec.left_margin = Inches(.77)
        sec.right_margin = Inches(.77)
        sec.header_distance = Inches(.35)
        sec.footer_distance = Inches(.38)
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(0,0,0)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(7)
        title = styles["Title"]
        title.font.name = "Times New Roman"
        title.font.size = Pt(16)
        title.font.bold = True
        title.font.color.rgb = RGBColor(0,0,0)
        title.paragraph_format.space_after = Pt(8)
        title_pr = title._element.get_or_add_pPr()
        title_border = title_pr.find(qn("w:pBdr"))
        if title_border is not None:
            title_pr.remove(title_border)
        for name, size in (("Heading 1", 12), ("Heading 2", 11)):
            s = styles[name]
            s.font.name = "Times New Roman"
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0,0,0)
            s.paragraph_format.space_before = Pt(13)
            s.paragraph_format.space_after = Pt(6)
            s.paragraph_format.keep_with_next = True
        header = sec.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(f"OBJECTIVE {self.objective}  |  RESEARCH RESULTS")
        run.font.name = "Times New Roman"; run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90,90,90)
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Page ")
        run.font.size = Pt(8)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)

    def para(self, text, lead=None):
        p = self.doc.add_paragraph()
        if lead:
            r = p.add_run(lead + " ")
            r.bold = True
        p.add_run(text)
        return p

    def caption(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True; r.font.size = Pt(10)
        return p

    def note(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run("Note. " + text)
        r.italic = True; r.font.size = Pt(8.5)
        return p

    def table(self, headers, lines, widths):
        tab = self.doc.add_table(rows=1, cols=len(headers))
        tab.autofit = False
        for i, width in enumerate(widths):
            tab.columns[i].width = Inches(width)
        for i, h in enumerate(headers):
            cell = tab.rows[0].cells[i]
            cell.width = Inches(widths[i])
            cell.text = h
            shade(cell, "424242");set_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cell, 90, 70, 90, 70)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = "Times New Roman";r.font.size=Pt(8);r.font.bold=True
                    r.font.color.rgb = RGBColor(255,255,255)
        repeat_header(tab.rows[0]);no_split(tab.rows[0])
        for j, line in enumerate(lines):
            row = tab.add_row();no_split(row)
            for i, value in enumerate(line):
                cell=row.cells[i]
                cell.width=Inches(widths[i]);cell.text=str(value)
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margin(cell);set_border(cell)
                if j%2: shade(cell,"F6F6F6")
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=1.0
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name="Times New Roman";r.font.size=Pt(8)
        self.doc.add_paragraph().paragraph_format.space_after=Pt(0)

    def interpretations(self, global_text, india_text, synthesis):
        self.doc.add_heading("Interpretation",level=2)
        self.para(global_text,"Global sample.")
        self.para(india_text,"India sample.")
        if not getattr(self, "figure_mode", False):
            self.para(synthesis,"Interpretive implication.")

    def top_model(self,sample,target):
        return max(pick(self.data[4],sample,target),key=lambda r: float(r["PR_AUC"]))

    def comparison(self,sample,target,model):
        return pick(self.data[6],sample,target,model)

    def add_intro(self):
        self.doc.add_paragraph(f"Objective {self.objective} results and interpretation",style="Title")
        p=self.doc.add_paragraph()
        p.paragraph_format.space_after=Pt(12)
        r=p.add_run(f"Socio-cognitive traits and the {THEMES[self.objective]}")
        r.italic=True;r.font.size=Pt(11)
        self.doc.add_heading("Results and analytical scope",level=1)
        self.para("This manuscript reports predictive analyses for total early-stage entrepreneurial activity (TEA), "
                  "opportunity-motivated entrepreneurship (OME), and necessity-motivated entrepreneurship (NME). "
                  "Each outcome was modeled separately for the global sample, which includes India, and for India alone. "
                  "The principal specification contains socio-cognitive and institutional predictors together with the "
                  "pre-specified control variables. Independent-variable-only results are retained for matched comparisons.")
        self.para("All estimates, figures, and tables were reconstructed from the completed saved model outputs. "
                  "The ten classifiers were not refitted. Precision-recall (PR) area and receiver operating characteristic (ROC) "
                  "area summarize discrimination; precision, recall, the F1 score, and confusion counts use thresholds "
                  "selected on validation data and evaluated on held-out test cases. The reported comparisons describe "
                  "predictive associations and do not identify causal effects.")

    def add_table1(self):
        self.doc.add_heading("Sample composition",level=1)
        self.caption("Table 1. Sample size, outcome prevalence, and analysis splits")
        d=self.data[1]
        self.table(["Sample","Outcome","Respondents","Positive cases","Prevalence","Test cases"],
                   [[r["Sample"],r["Target"],r["N"],r["Positive"],r["Prevalence"],r["Test"]] for r in d],
                   [1.05,.75,1.22,1.25,1.25,1.22])
        self.note("Global estimates include India. Training and validation counts are available in the accompanying table file; the same respondents enter paired predictor-set comparisons.")
        g=pick(d,"Global","TEA")[0];i=pick(d,"India","TEA")[0]
        gp=", ".join(f"{r['Target']} {r['Prevalence']}" for r in d if r["Sample"]=="Global")
        ip=", ".join(f"{r['Target']} {r['Prevalence']}" for r in d if r["Sample"]=="India")
        self.interpretations(
            f"The global analytical sample contains {g['N']} respondents, with {g['Test']} reserved for testing in each outcome analysis. Positive-case prevalence is {gp}. The lower prevalence of opportunity- and necessity-motivated entrepreneurship means that accuracy alone would provide an incomplete assessment of classification performance.",
            f"The India sample contains {i['N']} respondents and {i['Test']} held-out observations per outcome. Its observed prevalence is {ip}. The separate India models assess within-country prediction rather than the performance of a global model transported to India.",
            "Because the global sample includes India, the two samples are not independent. Differences in performance should therefore be interpreted as descriptive differences between fitted populations, with attention to the much smaller India test set.")

    def add_table2(self):
        self.doc.add_heading("Predictor distributions",level=1)
        self.caption("Table 2. Descriptive statistics and missingness of model predictors")
        d=self.data[2]
        self.table(["Predictor","Role","Global value","Missing","India value","Missing"],
                   [[pretty(r["Variable"]),"Predictor" if r["Role"]=="independent" else "Control",
                     r["Global"],r["Global_missing"],r["India"],r["India_missing"]] for r in d],
                   [1.92,.80,1.17,.70,1.17,.70])
        self.note("Binary values are proportions coded one; other numerical values are means with standard deviations in parentheses. Missing percentages refer to the eligible sample. Institutional scores repeat for respondents in the same country and survey year.")
        inst="prosperity_index_score" if self.objective==2 else "EFW"
        x=next(r for r in d if r["Variable"]==inst)
        self.interpretations(
            f"The global distribution combines respondent-level socio-cognitive variables with country-year institutional measures. The {pretty(inst).lower()} has a global summary of {x['Global']} and {x['Global_missing']} missingness. Several socio-cognitive fields have substantially more missing values than the institutional index, so the handling of missing categories is consequential for the fitted models.",
            f"In India, the same institutional measure has a summary of {x['India']} and {x['India_missing']} missingness. Within India, each respondent surveyed in a given year receives the same national institutional score. The respondent count should therefore not be interpreted as the number of independent observations of institutional conditions.",
            "The table describes observed predictor distributions, not adjusted relationships with entrepreneurship. Institutional score differences between India and the global sample may reflect country composition and survey-year coverage as well as substantive differences in the measured conditions.")

    def add_table3(self):
        self.doc.add_heading("Survey response pattern",level=1)
        self.caption("Table 3. Missing entrepreneurial intention by observed outcome")
        d=self.data[3]
        self.table(["Sample","Outcome","Positive cases missing","Negative cases missing"],
                   [[r["Sample"],r["Target"],r["Positive_pct"],r["Negative_pct"]] for r in d],
                   [1.10,.88,2.35,2.35])
        self.note("Entrepreneurial intention was retained as a predictor as specified in the analysis plan; its missing state was preserved rather than removed.")
        gn=", ".join(f"{r['Target']} {r['Negative_pct']}" for r in d if r["Sample"]=="Global")
        inn=", ".join(f"{r['Target']} {r['Negative_pct']}" for r in d if r["Sample"]=="India")
        self.interpretations(
            f"Entrepreneurial intention is missing for every positive case in all three global outcomes. Among negative cases, the corresponding missingness rates are {gn}. Missingness therefore carries substantial information about the observed outcome, independently of the respondent's stated intention value.",
            f"The India analysis exhibits the same complete missingness among positive cases. The negative-case rates are {inn}. This pattern applies to all India outcomes and is present in the predictor data used by the classifiers.",
            "The survey response pattern is essential to interpreting subsequent discrimination measures. High held-out performance reflects the entire observed-data process, including outcome-linked missingness, and should not be presented as evidence that institutional measures alone predict entrepreneurship with the reported accuracy.")

    def add_table4(self):
        self.doc.add_heading("Classifier discrimination",level=1)
        self.caption("Table 4. Highest precision-recall area model for each outcome and sample")
        d=[]
        for sample in ("Global","India"):
            for target in OUTCOMES:
                x=self.top_model(sample,target)
                d.append([sample,target,x["Model"],x["ROC_AUC"],x["PR_AUC"],x["Precision"],x["Recall"],x["F1"]])
        self.table(["Sample","Outcome","Selected model","ROC area","PR area","Precision","Recall","F1"],d,
                   [.68,.55,1.63,.78,.78,.77,.77,.78])
        self.note("The selected model has the largest held-out precision-recall area within each sample and outcome. All ten model results appear in Supplementary Table S1. Selection here is descriptive and does not imply a statistically significant advantage.")
        def selected(sample):
            return "; ".join(f"{t}: {self.top_model(sample,t)['Model']} "
                             f"(precision-recall area {self.top_model(sample,t)['PR_AUC']})"
                             for t in OUTCOMES)
        self.interpretations(
            f"The best observed precision-recall area in the global sample is {selected('Global')}. Discrimination is strongest for total early-stage entrepreneurial activity, while the two motive-specific outcomes are more difficult to identify because they have fewer positive cases and differ in their outcome composition.",
            f"For India, the corresponding selected classifiers are {selected('India')}. The leading algorithm differs across outcomes, indicating that the relative performance of model families depends on the target. Estimates are based on the smaller India test sample and should be read with this sampling difference in mind.",
            "The near-ceiling receiver operating characteristic areas must be interpreted alongside Table 3. They measure ranking on the saved held-out split under the observed survey missingness structure, not the predictive contribution of any one socio-cognitive or institutional variable.")

    def add_table5(self):
        self.doc.add_heading("Classification at selected thresholds",level=1)
        self.caption("Table 5. Confusion counts for the selected classifiers")
        d=[]
        for sample in ("Global","India"):
            for target in OUTCOMES:
                m=self.top_model(sample,target)["Model"]
                x=pick(self.data[5],sample,target,m)
                d.append([sample,target,m,x["TP"],x["FP"],x["TN"],x["FN"]])
        self.table(["Sample","Outcome","Selected model","True positive","False positive","True negative","False negative"],d,
                   [.71,.59,1.54,.97,.98,.98,.98])
        self.note("Each model's threshold was chosen using validation observations. Complete confusion matrices for all ten classifiers are provided in Supplementary Table S2.")
        def c(sample):
            out=[]
            for t in OUTCOMES:
                x=pick(self.data[5],sample,t,self.top_model(sample,t)["Model"])
                out.append(f"{t}: {x['FP']} false positives and {x['FN']} false negatives")
            return "; ".join(out)
        self.interpretations(
            f"At their validation-selected thresholds, the selected global classifiers produce {c('Global')}. These counts show the operating-point trade-off that is hidden by a single curve-area statistic. In particular, the smaller positive classes make false positive classifications consequential for precision.",
            f"For India, the selected models produce {c('India')}. The absolute counts are lower because the India test set is smaller. Model comparisons should use both these counts and the corresponding precision and recall values rather than the raw number of errors alone.",
            "No application-specific cost of a false positive or false negative was specified. Consequently, the reported validation-selected thresholds provide reproducible operating points but do not establish an optimal threshold for a policy or screening decision.")

    def add_table6(self):
        self.doc.add_heading("Incremental information from controls",level=1)
        self.caption("Table 6. Change in precision-recall area after including control variables")
        d=[]
        for sample in ("Global","India"):
            for target in OUTCOMES:
                m=self.top_model(sample,target)["Model"]
                x=self.comparison(sample,target,m)
                med=median(float(z["Delta_PR"]) for z in pick(self.data[6],sample,target))
                d.append([sample,target,m,x["PR_IV"],x["PR_IV_CV"],x["Delta_PR"],f"{med:+.3f}"])
        self.table(["Sample","Outcome","Selected model","Predictors only","With controls","Change","Median change all models"],d,
                   [.65,.56,1.40,1.02,1.02,.76,1.57])
        self.note("Change is the score with controls minus the predictor-only score on exactly the same held-out respondents. Full paired results are shown in Supplementary Table S3.")
        def summarize(sample):
            return "; ".join(f"{t} {self.comparison(sample,t,self.top_model(sample,t)['Model'])['Delta_PR']}"
                             for t in OUTCOMES)
        self.interpretations(
            f"For the selected global model in each outcome, the precision-recall area changes are {summarize('Global')}. The final column summarizes the corresponding change across all ten model families, guarding against an interpretation based only on the selected classifier.",
            f"The selected India model changes are {summarize('India')}. Survey year is included in the control specification. Since India's institutional scores are annual country-level values, year and the institutional indices capture overlapping temporal variation.",
            "These matched test-split differences describe incremental predictive information from the full control set. They do not estimate the independent effect of age, survey year, macroeconomic conditions, or any single institutional pillar.")

    def add_table7(self):
        self.doc.add_page_break()
        self.doc.add_heading("Institutional predictors",level=1)
        self.caption("Table 7. Leading institutional predictors in the random forest")
        d=[]
        for sample in ("Global","India"):
            for target in OUTCOMES:
                x=pick(self.data[7],sample,target)[0]
                d.append([sample,target,pretty(x["Variable"]),x["Share"],x["Institution_total"]])
        self.table(["Sample","Outcome","Leading institutional variable","Share of all importance","Institutional total"],d,
                   [.72,.58,2.30,1.35,1.68])
        self.note("Importance is the random forest's impurity-based allocation across fitted inputs. The institutional total sums all raw institutional fields, while the leading-variable column reports the largest one. Missingness indicators enter the denominator but are excluded from the institutional ranking.")
        def summary(sample):
            return "; ".join(f"{t}: {pretty(pick(self.data[7],sample,t)[0]['Variable'])} "
                             f"({pick(self.data[7],sample,t)[0]['Share']} of all importance)"
                             for t in OUTCOMES)
        self.interpretations(
            f"The most prominent institutional dimensions in the global fits are {summary('Global')}. The reported shares indicate how often the model used these fields in impurity-reducing splits relative to all predictors, including socio-cognitive variables and controls.",
            f"In the India fits, the leading institutional fields are {summary('India')}. These scores vary at the national annual level, yielding only eight distinct survey-year profiles. A large or small importance share is therefore sensitive to the particular temporal pattern and to competition with the year control.",
            "Random-forest impurity importance is not an effect estimate and may redistribute importance among correlated institutional pillars. The table supports a descriptive account of model reliance, not a ranking of which policy reform would change entrepreneurship rates.")

    def add_figures(self):
        self.doc.add_heading("Figures",level=1)
        titles={1:"Receiver operating characteristic curves",2:"Precision recall curves",
                3:"Random forest feature importance",4:"Change in precision recall area after adding controls",
                5:"Institutional feature importance"}
        notes={
            1:"The five largest receiver operating characteristic areas are plotted within each sample and outcome panel. The dotted diagonal represents chance ranking.",
            2:"The five largest precision-recall areas are plotted within each panel. The dotted horizontal line marks outcome prevalence.",
            3:"Bars show the five most important inputs in each random forest, expressed as a share of total impurity importance.",
            4:"Dots show the change in held-out precision-recall area for each of the ten model families when controls are included.",
            5:"Bars show the five most important raw institutional variables in each random forest, as a share of total model importance.",
        }
        for number in range(1,6):
            self.doc.add_page_break()
            self.caption(f"Figure {number}. {titles[number]}")
            p=self.doc.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(2)
            figure_width = 5.0 if self.objective == 3 and number == 5 else 5.45
            p.add_run().add_picture(str(self.dir/"word_figures"/f"figure_{number}.png"),width=Inches(figure_width))
            self.note(notes[number])
            self._figure_interpretation(number)

    def _figure_interpretation(self,n):
        self.figure_mode = True
        if n in (1,2):
            metric="ROC_AUC" if n==1 else "PR_AUC"
            metric_name="receiver operating characteristic area" if n==1 else "precision-recall area"
            def top(sample):
                return "; ".join(f"{t} {self.top_model(sample,t)['Model']} "
                                 f"({self.top_model(sample,t)[metric]})" for t in OUTCOMES)
            self.interpretations(
                f"The global panels show consistently strong ranking across the three outcomes. For the models selected by precision-recall area, the corresponding {metric_name} values are {top('Global')}. The spread between curves is much smaller than the separation from the reference line, indicating that several model families perform similarly on the saved test split.",
                f"The India panels show {metric_name} values of {top('India')} for the selected classifiers. Differences in curve shape should be read against the smaller India test sample and the prevalence of each positive outcome. The India curves describe separately trained country-specific fits.",
                ("Receiver operating characteristic curves provide a useful ranking comparison but can appear strong even when positive cases are scarce. The precision-recall results and confusion counts give the complementary view of positive-case identification."
                 if n==1 else
                 "The prevalence reference lines show why a common precision benchmark cannot be used across the three outcomes. Curves above their respective baselines indicate better-than-prevalence ranking, while the decline in precision at higher recall represents the cost of detecting more positive cases."))
        elif n==3:
            def top(sample):
                # The complete saved importance files are used for the figure.
                import csv as _csv
                batch=HERE.parents[1]/"analysis_output"/f"objective{self.objective}_batch_20260928"
                bits=[]
                for target in OUTCOMES:
                    file=batch/f"objective_{self.objective}"/sample.lower()/"iv_cv"/target/"random_forest_importance.csv"
                    data=rows(file)
                    x=max(data,key=lambda r:float(r["importance"]))
                    bits.append(f"{target}: {pretty(x['variable'])}")
                return "; ".join(bits)
            self.interpretations(
                f"The leading inputs in the global panels are {top('Global')}. The figure also shows that the importance profile changes across the three outcome definitions. Feature importance here records the fitted forest's use of a variable in splitting the data, rather than the direction of its relationship with entrepreneurship.",
                f"The India panels identify {top('India')} as the leading inputs. The narrower country-specific sample and limited variation in national institutional scores can change the ordering relative to the global fits.",
                "The survey pattern documented in Table 3 is relevant when entrepreneurial intention appears among the leading predictors. A missing-state contribution should not be described as evidence that stated intention is an unusually strong behavioral precursor in this dataset.")
        elif n==4:
            def summary(sample):
                parts=[]
                for target in OUTCOMES:
                    group=pick(self.data[6],sample,target)
                    d=[float(r["Delta_PR"]) for r in group]
                    parts.append(f"{target}: {sum(v>0 for v in d)} of ten models improve; median change {median(d):+.3f}")
                return "; ".join(parts)
            self.interpretations(
                f"In the global analysis, {summary('Global')}. The horizontal distance from zero makes the incremental predictive value of the entire control set visible for each classifier and outcome.",
                f"In India, {summary('India')}. The spread across model families indicates that adding controls does not contribute a fixed amount of information independent of the learning algorithm.",
                "All pairs share the same held-out respondents. The chart therefore isolates a change in the predictor specification on this split, but it does not isolate the contribution of a particular control or establish a causal effect.")
        else:
            def summary(sample):
                return "; ".join(f"{t}: {pretty(pick(self.data[7],sample,t)[0]['Variable'])} "
                                 f"({pick(self.data[7],sample,t)[0]['Share']})" for t in OUTCOMES)
            self.interpretations(
                f"The leading institutional fields in the global panels are {summary('Global')}. Their bars use all fitted predictors as the denominator, allowing the reader to judge institutional usage in the context of the wider model rather than only relative to other institutional pillars.",
                f"For India, the corresponding leading fields are {summary('India')}. Because the scores are shared by respondents in the same year, the bars primarily summarize how the model uses annual national variation.",
                "The ordering is descriptive and can be altered by correlations among index components, missingness indicators, and survey year. It does not show that changing an institutional score would cause a proportional change in entrepreneurial activity.")
        self.figure_mode = False

    def supplement(self):
        self.doc.add_page_break()
        self.doc.add_heading("Supplementary model tables",level=1)
        self.para("The following tables retain all ten fitted model families for each outcome and sample. They are included for auditability; the principal tables in the results section summarize the selected models.")
        self.caption("Supplementary Table S1. Complete held-out classifier performance")
        d=self.data[4]
        self.table(["Sample","Outcome","Model","ROC area","PR area","Precision","Recall","F1"],
                   [[r["Sample"],r["Target"],r["Model"],r["ROC_AUC"],r["PR_AUC"],r["Precision"],r["Recall"],r["F1"]] for r in d],
                   [.65,.53,1.55,.81,.81,.8,.8,.8])
        self.note("All figures are measured on held-out test observations using the full predictor and control specification.")
        self.para("Across the global outcomes, model rankings differ in detail but several nonlinear models achieve similar discrimination. The complete table prevents the six selected results in Table 4 from obscuring the performance of alternative model families.","Global sample.")
        self.para("The India rankings also vary by outcome. Comparison of the full rows is preferable to treating a small difference between the two leading models as conclusive, particularly with the smaller country-specific test sample.","India sample.")
        self.caption("Supplementary Table S2. Complete held-out confusion matrices")
        d=self.data[5]
        self.table(["Sample","Outcome","Model","True positive","False positive","True negative","False negative"],
                   [[r["Sample"],r["Target"],r["Model"],r["TP"],r["FP"],r["TN"],r["FN"]] for r in d],
                   [.69,.57,1.52,.99,.99,.99,.99])
        self.note("All confusion counts use the validation-selected threshold for the indicated classifier.")
        self.para("The global counts show how model families with similar ranking measures can produce different threshold classifications. This distinction is particularly relevant for the less prevalent motivated-entrepreneurship outcomes.","Global sample.")
        self.para("India confusion counts are smaller in absolute terms because fewer observations are tested. Their interpretation depends on the balance between false positive and false negative decisions rather than on total errors alone.","India sample.")
        self.caption("Supplementary Table S3. Complete matched predictor-set comparison")
        self.para("The global model-level changes show that the value of the additional control variables is heterogeneous across algorithms and outcomes. These are observed differences in predictive ranking rather than statistical estimates of a structural parameter.","Global sample.")
        self.para("India changes should also be interpreted in light of the temporal overlap between survey year and national institutional indices. The table quantifies the predictive change from the whole control set and does not separate a year effect from an institutional effect.","India sample.")
        d=self.data[6]
        self.table(["Sample","Outcome","Model","PR predictors","PR with controls","Change in PR area"],
                   [[r["Sample"],r["Target"],r["Model"],r["PR_IV"],r["PR_IV_CV"],r["Delta_PR"]] for r in d],
                   [.72,.58,1.72,1.19,1.19,1.34])
        self.note("Positive change means that the specification with controls has a larger held-out precision-recall area. Every pair uses identical test records.")

    def build(self):
        self.add_intro()
        for method in (self.add_table1,self.add_table2,self.add_table3,self.add_table4,
                       self.add_table5,self.add_table6,self.add_table7,self.add_figures,
                       self.supplement):
            method()
        path=self.dir/f"Objective{self.objective}_Journal_Style_Results.docx"
        self.doc.save(path)
        print(path)


if __name__=="__main__":
    for objective in (2,3): Manuscript(objective).build()
